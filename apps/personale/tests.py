import os
from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth.models import Group, User
from django.test import TestCase

from . import calcolo, importa_excel
from .models import Dipendente, MesePersonale, ParametroRetributivo, Presenza, RigaMese

OTTOBRE = date(2026, 10, 1)
# Excel reali (solo in locale, fuori dal repository): percorsi da variabili d'ambiente
FILE_ELENCO = os.environ.get('PERSONALE_EXCEL_ELENCO', '')
FILE_MESE = os.environ.get('PERSONALE_EXCEL_MESE', '')


def utente(nome, gruppo):
    u = User.objects.create_user(nome, password='x')
    u.groups.add(Group.objects.get_or_create(name=gruppo)[0])
    return u


def dipendente(nome, tariffa='6', coeff='1.5', tipo='contrattualizzato', dal=date(2025, 10, 1)):
    d = Dipendente.objects.create(nome=nome, tipo=tipo)
    ParametroRetributivo.objects.create(dipendente=d, valido_dal=dal, tariffa_oraria=Decimal(tariffa),
                                        coefficiente_lordo=Decimal(coeff))
    return d


def presenze(dip, valori, mese=OTTOBRE):
    for i, v in enumerate(valori):
        ore, assenza = calcolo.leggi_cella(v)
        if ore is not None or assenza:
            Presenza.objects.create(dipendente=dip, data=mese + timedelta(days=i), ore=ore, assenza=assenza)


class CalcoloTest(TestCase):
    def setUp(self):
        # dati inventati: 20 giorni da 8,5 ore, 6 di ferie, 3 liberi
        self.mario = dipendente('Mario')
        presenze(self.mario, ['8.5'] * 10 + ['F'] * 6 + ['L'] * 3 + ['8.5'] * 10)
        self.mese = MesePersonale.objects.create(mese=OTTOBRE)
        self.riga = RigaMese.objects.create(mese=self.mese, dipendente=self.mario, busta_paga=Decimal('1400'))

    def test_riga_come_excel(self):
        c = calcolo.calcola_mese(self.mese)
        r = c['righe'][0]
        self.assertEqual(r['ore_lavorate'], Decimal('170'))
        self.assertEqual((r['gg_ferie'], r['gg_libero']), (6, 3))
        self.assertEqual(r['ore_totali'], Decimal('218'))
        self.assertEqual(r['netto'], Decimal('1308'))
        self.assertEqual(r['rateo'], Decimal('109'))
        self.assertEqual(r['lordo'], Decimal('1962'))
        self.assertEqual(r['contributi'], Decimal('700'))
        self.assertEqual(r['differenza'], Decimal('17'))
        self.assertEqual(r['costo'], Decimal('2117'))
        self.assertEqual(c['riepilogo']['attivi'], 1)
        # netto incassato = netto + 13a + bonus = busta paga + fuori busta
        self.assertEqual(r['netto_incassato'], Decimal('1417'))
        self.assertEqual(r['netto_incassato'], r['busta_paga'] + r['totale_pagare'])
        # indice di costo orario: costo aziendale / ore lavorate (non le ore pagate di ferie)
        self.assertEqual(c['riepilogo']['costo_ora'], Decimal('2117') / Decimal('170'))
        self.assertEqual(c['da_verificare'], 0)

    def test_bonus_pagato_e_occasionale(self):
        self.riga.bonus, self.riga.pagato = Decimal('50'), Decimal('20')
        self.riga.save()
        r = calcolo.calcola_mese(self.mese)['righe'][0]
        self.assertEqual(r['totale_pagare'], Decimal('67'))
        self.assertEqual(r['residuo'], Decimal('47'))
        self.assertEqual(r['costo'], Decimal('2167'))
        self.assertEqual(r['netto_incassato'], Decimal('1467'))
        self.mario.tipo = 'occasionale'
        self.mario.save()
        self.assertEqual(calcolo.calcola_mese(self.mese)['righe'][0]['rateo'], 0)

    def test_parametri_per_data_e_coefficiente_del_mese(self):
        ParametroRetributivo.objects.create(dipendente=self.mario, valido_dal=date(2026, 11, 1),
                                            tariffa_oraria=Decimal('7'), coefficiente_lordo=Decimal('1.2'))
        self.assertEqual(calcolo.parametri_per(self.mario, OTTOBRE).tariffa_oraria, Decimal('6'))
        self.assertEqual(calcolo.parametri_per(self.mario, date(2026, 11, 1)).tariffa_oraria, Decimal('7'))
        self.assertIsNone(calcolo.parametri_per(self.mario, date(2025, 9, 1)))
        self.riga.coefficiente_lordo = Decimal('1.17')
        self.riga.save()
        r = calcolo.calcola_mese(self.mese)['righe'][0]
        self.assertTrue(r['coeff_sovrascritto'])
        self.assertEqual(r['contributi'], Decimal('238'))

    def test_controlli(self):
        nuovo = Dipendente.objects.create(nome='Nuovo')
        presenze(nuovo, ['8'])
        RigaMese.objects.create(mese=self.mese, dipendente=nuovo)
        self.riga.busta_paga = None
        self.riga.save()
        controlli = {k['nome']: k for k in calcolo.calcola_mese(self.mese)['controlli']}
        self.assertFalse(controlli['Parametri validi per il mese']['ok'])
        self.assertIn('Nuovo', controlli['Tariffa inserita']['messaggio'])
        self.assertIn('Mario', controlli['Busta paga inserita']['messaggio'])

    def test_celle(self):
        self.assertEqual(calcolo.leggi_cella('8,5'), (Decimal('8.50'), ''))
        self.assertEqual(calcolo.leggi_cella(' f '), (None, 'F'))
        self.assertEqual(calcolo.leggi_cella(''), (None, ''))
        for errato in ('17', 'X', '-1'):
            with self.assertRaises(ValueError):
                calcolo.leggi_cella(errato)
        self.assertEqual(calcolo.testo_numero(Decimal('10.00')), '10')
        self.assertEqual(calcolo.testo_numero(Decimal('1.170')), '1,17')

    def test_andamento_anno(self):
        a = calcolo.andamento_anno(2026)
        self.assertEqual(a['totale']['mesi'], 1)
        self.assertEqual(a['mesi'][9]['costo'], Decimal('2117'))
        self.assertEqual(a['totale']['costo_ora'], Decimal('2117') / Decimal('170'))
        self.assertEqual(a['trimestri'][3]['mesi'], 1)
        self.assertEqual(a['per_dipendente'][0]['dipendente'], self.mario)


class VisteTest(TestCase):
    def setUp(self):
        from apps.auth_system.models import PermessoGruppo
        self.titolare = utente('titolare1', 'titolare')
        PermessoGruppo.objects.get_or_create(sezione='personale', gruppo=Group.objects.get(name='titolare'))
        self.mario = dipendente('Mario')
        self.client.force_login(self.titolare)

    def test_operatore_senza_permesso(self):
        self.client.force_login(utente('op1', 'operatore'))
        risposta = self.client.post('/personale/', {'azione': 'crea'}, HTTP_ACCEPT='application/json')
        self.assertEqual(risposta.status_code, 403)
        self.assertFalse(MesePersonale.objects.exists())

    def test_crea_e_salva_mese(self):
        self.client.post('/personale/?mese=2026-10', {'azione': 'crea'})
        riga = RigaMese.objects.get(mese__mese=OTTOBRE)
        dati = {'azione': 'salva', 'ore_giorno_ferie': '8', 'ore_giorno_malattia': '8',
                f'c_{riga.pk}_1': '8,5', f'c_{riga.pk}_2': 'f', f'c_{riga.pk}_3': 'L',
                f'coeff_{riga.pk}': '', f'busta_{riga.pk}': '1400', f'bonus_{riga.pk}': '-10', f'pagato_{riga.pk}': ''}
        risposta = self.client.post('/personale/?mese=2026-10', dati)
        self.assertEqual(risposta.status_code, 302)
        self.assertEqual(Presenza.objects.get(data=OTTOBRE).ore, Decimal('8.5'))
        self.assertEqual(Presenza.objects.get(data=date(2026, 10, 2)).assenza, 'F')
        riga.refresh_from_db()
        self.assertEqual((riga.busta_paga, riga.bonus, riga.coefficiente_lordo), (Decimal('1400'), Decimal('-10'), None))

        # una cella vuota toglie la presenza
        dati[f'c_{riga.pk}_1'] = ''
        self.client.post('/personale/?mese=2026-10', dati)
        self.assertFalse(Presenza.objects.filter(data=OTTOBRE).exists())

    def test_cella_errata_non_salva_nulla(self):
        self.client.post('/personale/?mese=2026-10', {'azione': 'crea'})
        riga = RigaMese.objects.get(mese__mese=OTTOBRE)
        from django.test import RequestFactory
        from django.contrib.messages.storage.fallback import FallbackStorage
        from .views import mese
        req = RequestFactory().post('/personale/?mese=2026-10', {
            'azione': 'salva', f'c_{riga.pk}_1': '8', f'c_{riga.pk}_2': '17', f'busta_{riga.pk}': '1000'})
        req.user, req.session = self.titolare, {}
        req._messages = FallbackStorage(req)
        risposta = mese(req)
        self.assertEqual(risposta.status_code, 200)
        self.assertFalse(Presenza.objects.exists())
        riga.refresh_from_db()
        self.assertIsNone(riga.busta_paga)
        self.assertIn('value="17"', risposta.content.decode())

    def test_parametri_e_dipendenti(self):
        self.client.post('/personale/dipendenti/', {'azione': 'salva_dipendente', 'nome': 'Anna',
                                                    'tipo': 'occasionale', 'attivo': '1'})
        k = Dipendente.objects.get(nome='Anna')
        self.assertEqual(k.tipo, 'occasionale')
        self.client.post('/personale/dipendenti/', {'azione': 'aggiungi_parametro', 'dipendente': k.pk,
                                                    'valido_dal': '2026-10-01', 'tariffa_oraria': '6,5',
                                                    'coefficiente_lordo': '1,17'})
        p = k.parametri.get()
        self.assertEqual((p.tariffa_oraria, p.coefficiente_lordo), (Decimal('6.5'), Decimal('1.17')))
        # stesso "valido dal": rifiutato, nessun doppione
        self.client.post('/personale/dipendenti/', {'azione': 'aggiungi_parametro', 'dipendente': k.pk,
                                                    'valido_dal': '2026-10-01', 'tariffa_oraria': '7'})
        self.assertEqual(k.parametri.count(), 1)
        # nome gia' usato
        self.client.post('/personale/dipendenti/', {'azione': 'salva_dipendente', 'nome': 'mario'})
        self.assertEqual(Dipendente.objects.filter(nome__iexact='mario').count(), 1)


class ImportExcelTest(TestCase):
    """Con i due Excel reali, indicati da PERSONALE_EXCEL_ELENCO e
    PERSONALE_EXCEL_MESE (senza si salta): il calcolo deve dare gli stessi
    valori che Excel ha salvato nel foglio del mese."""

    def setUp(self):
        if not (os.path.isfile(FILE_ELENCO) and os.path.isfile(FILE_MESE)):
            self.skipTest('File Excel non indicati')

    def _valori_excel(self):
        """Valori del calcolo per dipendente e totali, come li ha salvati Excel."""
        import openpyxl
        ws = openpyxl.load_workbook(FILE_MESE, data_only=True).active
        r_calc = next(r for r in range(1, ws.max_row + 1)
                      if str(ws.cell(r, 1).value or '').startswith('Calcolo retribuzione'))
        col = {str(ws.cell(r_calc + 1, c).value or ''): c for c in range(1, ws.max_column + 1)}
        chiavi = {'ore': 'Ore totali', 'netto': 'Netto (ore × tariffa)', 'lordo': 'Lordo',
                  'costo': 'Costo aziendale', 'coeff': 'Coeff. lordo'}
        righe = {}
        for r in range(r_calc + 2, ws.max_row + 1):
            nome = str(ws.cell(r, 1).value or '')
            valori = {k: Decimal(str(ws.cell(r, col[t]).value or 0)) for k, t in chiavi.items()}
            if nome == 'Totale':
                return righe, valori
            if nome:
                righe[nome] = valori
        raise AssertionError('Riga Totale non trovata')

    def test_elenco_e_mese_come_excel(self):
        with open(FILE_ELENCO, 'rb') as f:
            self.assertEqual(importa_excel.importa(f)['tipo'], 'elenco')
        self.assertTrue(Dipendente.objects.exists())
        with open(FILE_MESE, 'rb') as f:
            esito = importa_excel.importa(f)
        self.assertEqual(esito['tipo'], 'mese')
        self.assertEqual(esito['avvisi'], [])

        righe_excel, totali_excel = self._valori_excel()
        mese_obj = MesePersonale.objects.get(mese=esito['mese'])
        c = calcolo.calcola_mese(mese_obj)
        # Excel salva numeri in virgola mobile (1560,2250000000001): mezzo centesimo di tolleranza
        uguale = lambda a, b, msg: self.assertLessEqual(abs(a - b), Decimal('0.005'), f'{msg}: {a} != {b}')
        self.assertEqual(len(c['righe']), len(righe_excel))
        self.assertEqual(c['totali']['ore_totali'], totali_excel['ore'])
        for k in ('netto', 'lordo', 'costo'):
            uguale(c['totali'][k], totali_excel[k], k)
        for r in c['righe']:
            excel = righe_excel[r['dipendente'].nome]
            uguale(r['costo'], excel['costo'], r['dipendente'].nome)
            # il coefficiente usato nel mese e' quello del foglio, anche se diverso dall'elenco
            self.assertEqual(r['coeff'], excel['coeff'], r['dipendente'].nome)

        # reimportare non duplica nulla e da' gli stessi conti
        with open(FILE_MESE, 'rb') as f:
            importa_excel.importa(f)
        self.assertEqual(RigaMese.objects.count(), len(righe_excel))
        uguale(calcolo.calcola_mese(mese_obj)['totali']['costo'], totali_excel['costo'], 'costo dopo reimport')

    def test_mese_senza_elenco(self):
        with open(FILE_MESE, 'rb') as f:
            esito = importa_excel.importa(f)
        self.assertIn("Non nell'elenco", esito['avvisi'][0])

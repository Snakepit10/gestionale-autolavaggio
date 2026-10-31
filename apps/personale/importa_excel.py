"""Import dai fogli Excel usati finora (riconosciuti da soli):

- "Elenco dipendenti" (file Retribuzione): tabelle Dipendenti (Nome,
  Tipo, Note) e Parametri (Dipendente, Valido dal, Tariffa oraria,
  Coefficiente lordo, Note) -> Dipendente e ParametroRetributivo;
- "Registrazione mensile" (A3 = "Mese (1° giorno)", mese in B3):
  calendario presenze, busta paga, bonus, pagato e coefficiente del
  mese (sovrascrittura solo se diverso dai parametri).

Idempotente: reimportare lo stesso file non duplica nulla e il mese
prende esattamente i valori del file. I valori delle formule si leggono
dalla copia salvata da Excel (data_only)."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction

from .calcolo import giorni_del_mese, leggi_cella, parametri_per
from .models import Dipendente, MesePersonale, ParametroRetributivo, Presenza, RigaMese


class ErroreImport(Exception):
    pass


def _testo(v):
    return str(v).strip() if v is not None else ''


def _decimale(v):
    if v is None or _testo(v) == '':
        return None
    try:
        return Decimal(str(v).replace(',', '.')).quantize(Decimal('0.001'))
    except InvalidOperation:
        return None


def _data(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def importa(file):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(file, data_only=True)
    except Exception as e:                       # noqa: BLE001 - file non Excel o rovinato
        raise ErroreImport(f'Il file non si apre come Excel (.xlsx): {e}')
    if 'Elenco dipendenti' in wb.sheetnames:
        return importa_elenco(wb['Elenco dipendenti'])
    for ws in wb.worksheets:
        if _testo(ws['A3'].value).lower().startswith('mese (1'):
            return importa_mese(ws)
    raise ErroreImport('Formato non riconosciuto: serve il foglio "Elenco dipendenti" '
                       'oppure un foglio mensile con "Mese (1° giorno)" in A3.')


def _righe_tabella(ws, nome, intestazione):
    """Righe (dict per intestazione) della tabella Excel `nome`; se il
    foglio non ha tabelle cerca la riga d'intestazione."""
    if nome in ws.tables:
        celle = ws[ws.tables[nome].ref]
    else:
        celle = None
        for riga in ws.iter_rows():
            for i, c in enumerate(riga):
                if _testo(c.value) == intestazione[0] and all(
                        _testo(riga[i + j].value) == h for j, h in enumerate(intestazione) if i + j < len(riga)):
                    celle = [r[i:i + len(intestazione)] for r in ws.iter_rows(min_row=c.row)]
                    break
            if celle:
                break
        if not celle:
            raise ErroreImport(f'Tabella {nome} non trovata nel foglio "{ws.title}".')
    titoli = [_testo(c.value) for c in celle[0]]
    righe = []
    for r in celle[1:]:
        valori = dict(zip(titoli, (c.value for c in r)))
        if not _testo(valori.get(intestazione[0])):
            if nome in ws.tables:
                continue
            break
        righe.append(valori)
    return righe


@transaction.atomic
def importa_elenco(ws):
    nuovi = aggiornati = parametri_nuovi = 0
    for i, r in enumerate(_righe_tabella(ws, 'Dipendenti', ['Nome', 'Tipo'])):
        nome = _testo(r.get('Nome'))
        note = _testo(r.get('Note'))
        tipo = 'occasionale' if _testo(r.get('Tipo')).lower().startswith('occas') else 'contrattualizzato'
        _, creato = Dipendente.objects.update_or_create(
            nome__iexact=nome,
            defaults={'nome': nome, 'tipo': tipo, 'note': note, 'ordine': i,
                      'attivo': 'non più attivo' not in note.lower() and 'non piu attivo' not in note.lower()})
        nuovi += creato
        aggiornati += not creato

    avvisi = []
    for r in _righe_tabella(ws, 'Parametri', ['Dipendente', 'Valido dal']):
        nome = _testo(r.get('Dipendente'))
        dip = Dipendente.objects.filter(nome__iexact=nome).first()
        dal = _data(r.get('Valido dal'))
        tariffa = _decimale(r.get('Tariffa oraria'))
        if dip is None or dal is None or tariffa is None:
            avvisi.append(f'Parametri di "{nome}" saltati: dipendente, data o tariffa mancanti.')
            continue
        _, creato = ParametroRetributivo.objects.update_or_create(
            dipendente=dip, valido_dal=dal,
            defaults={'tariffa_oraria': tariffa,
                      'coefficiente_lordo': _decimale(r.get('Coefficiente lordo')) or Decimal('1'),
                      'note': _testo(r.get('Note'))})
        parametri_nuovi += creato
    return {
        'tipo': 'elenco',
        'messaggio': (f'Elenco dipendenti importato: {nuovi} nuovi, {aggiornati} aggiornati, '
                      f'{parametri_nuovi} nuove righe di parametri.'),
        'avvisi': avvisi,
    }


def _trova_riga(ws, testo, colonna=1, dopo=1):
    for r in range(dopo, ws.max_row + 1):
        if _testo(ws.cell(r, colonna).value).lower() == testo.lower():
            return r
    return None


@transaction.atomic
def importa_mese(ws):
    mese = _data(ws['B3'].value)
    if mese is None:
        raise ErroreImport('In B3 manca il mese (1° giorno).')
    mese = mese.replace(day=1)
    giorni = giorni_del_mese(mese)
    r_giorno = _trova_riga(ws, 'Giorno')
    r_totale = _trova_riga(ws, 'Totale ore giorno')
    r_calcolo = _trova_riga(ws, 'Calcolo retribuzione e costo aziendale')
    if not (r_giorno and r_totale and r_calcolo):
        raise ErroreImport('Foglio mensile diverso dal modello: non trovo calendario o calcolo.')

    mese_obj, _ = MesePersonale.objects.get_or_create(mese=mese)
    # ore pagate per giorno di ferie e malattia (colonna "Ore pagate/giorno")
    for r in range(r_giorno, r_giorno + 4):
        for c in range(1, ws.max_column + 1):
            if _testo(ws.cell(r, c).value).lower().startswith('f = ferie') and _decimale(ws.cell(r, c + 1).value):
                mese_obj.ore_giorno_ferie = _decimale(ws.cell(r, c + 1).value)
            if _testo(ws.cell(r, c).value).lower().startswith('m = malattia') and _decimale(ws.cell(r, c + 1).value):
                mese_obj.ore_giorno_malattia = _decimale(ws.cell(r, c + 1).value)
    mese_obj.save()

    # valori del calcolo per nome (celle gialle + coefficiente usato)
    col = {}
    r_intest = r_calcolo + 1
    for c in range(1, ws.max_column + 1):
        col[_testo(ws.cell(r_intest, c).value)] = c
    calcolo = {}
    for r in range(r_intest + 1, ws.max_row + 1):
        nome = _testo(ws.cell(r, 1).value)
        if nome.lower() == 'totale':
            break
        if nome:
            calcolo[nome.lower()] = {k: ws.cell(r, col[k]).value if k in col else None
                                     for k in ('Coeff. lordo', 'Busta paga', 'Bonus', 'Pagato')}

    mancanti, errori, importati = [], [], 0
    for ordine, r in enumerate(range(r_giorno + 2, r_totale)):
        nome = _testo(ws.cell(r, 1).value)
        if not nome:
            continue
        dip = Dipendente.objects.filter(nome__iexact=nome).first()
        if dip is None:
            mancanti.append(nome)
            continue
        valori = calcolo.get(nome.lower(), {})
        coeff = _decimale(valori.get('Coeff. lordo'))
        param = parametri_per(dip, mese)
        RigaMese.objects.update_or_create(mese=mese_obj, dipendente=dip, defaults={
            'ordine': ordine,
            'coefficiente_lordo': coeff if coeff and (param is None or coeff != param.coefficiente_lordo) else None,
            'busta_paga': _decimale(valori.get('Busta paga')),
            'bonus': _decimale(valori.get('Bonus')),
            'pagato': _decimale(valori.get('Pagato')),
        })
        for i, giorno in enumerate(giorni):
            try:
                ore, assenza = leggi_cella(ws.cell(r, 2 + i).value)
            except ValueError as e:
                errori.append(f'{nome} {giorno:%d/%m}: {e}')
                continue
            if ore is None and not assenza:
                Presenza.objects.filter(dipendente=dip, data=giorno, origine='manuale').delete()
            else:
                Presenza.objects.update_or_create(dipendente=dip, data=giorno, defaults={
                    'ore': ore, 'assenza': assenza, 'origine': 'manuale'})
        importati += 1

    avvisi = []
    if mancanti:
        avvisi.append('Non nell\'elenco dipendenti (importa prima l\'elenco): ' + ', '.join(mancanti))
    avvisi += errori
    return {
        'tipo': 'mese', 'mese': mese,
        'messaggio': f'Mese {mese:%m/%Y} importato: {importati} dipendenti.',
        'avvisi': avvisi,
    }
